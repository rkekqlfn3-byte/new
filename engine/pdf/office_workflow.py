"""Verified Office artifacts created from page-grounded PDF analysis."""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Mapping

from engine.app_actions.com_lifecycle import com_apartment
from engine.app_actions.excel_adapter import create_owned_excel_application
from engine.app_actions.value_normalizer import (
    excel_values_equal,
    normalize_excel_input,
)
from engine.pdf.analysis import PdfTableCandidate, citation_label
from engine.pdf.context import PdfConnection
from engine.workflow_step_registry import (
    pdf_workflow_step_recipe,
    validate_pdf_workflow_step_recipe,
)
from engine.workflows.business_workflow import (
    HwpReportWriter,
    PowerPointSummaryWriter,
    WordReportWriter,
    WorkflowError,
    WorkProductData,
    file_fingerprint,
)

_OUTPUT_SUFFIXES = {
    "excel": ".xlsx",
    "word": ".docx",
    "hwp": ".hwp",
    "powerpoint": ".pptx",
}


def _safe_stem(value: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "")).strip(" .")
    return (text[:80] or "PDF")


def _reserve_output(directory: Path, stem: str, suffix: str) -> Path:
    candidate = directory / f"{stem}{suffix}"
    if not candidate.exists():
        return candidate
    for index in range(2, 1_000):
        candidate = directory / f"{stem}_{index}{suffix}"
        if not candidate.exists():
            return candidate
    raise WorkflowError("PDF 결과물 파일 이름을 안전하게 예약하지 못했습니다.")


def _excel_cell_equal(actual, expected) -> bool:
    normalized = normalize_excel_input(expected).get("value")
    return excel_values_equal(actual, normalized)


class PdfExcelTableWriter:
    """Create a new XLSX and verify its dimensions and representative cells."""

    def __init__(self, application_factory=None, com_runtime=None):
        self._application_factory = application_factory
        self._com_runtime = com_runtime

    def run(self, output_path, table: PdfTableCandidate) -> dict[str, Any]:
        if not isinstance(table, PdfTableCandidate):
            raise WorkflowError("Excel 생성에는 검증된 PDF 표가 필요합니다.")
        destination = Path(output_path)
        if destination.exists() or destination.suffix.casefold() != ".xlsx":
            raise WorkflowError("새 Excel 결과물 경로가 올바르지 않습니다.")
        matrix = (tuple(table.headers), *tuple(table.rows))
        row_count = len(matrix)
        column_count = len(table.headers)
        lease = None
        created = False
        with com_apartment(self._com_runtime):
            try:
                lease = create_owned_excel_application(self._application_factory)
                application = lease.application
                application.Visible = False
                application.DisplayAlerts = False
                workbook = lease.register_owned_document(application.Workbooks.Add())
                worksheet = workbook.Worksheets(1)
                worksheet.Name = "PDF_Table"
                target = worksheet.Range(
                    worksheet.Cells(1, 1),
                    worksheet.Cells(row_count, column_count),
                )
                target.Value = matrix
                workbook.SaveAs(str(destination), FileFormat=51)
                created = destination.is_file()
                if not created:
                    raise WorkflowError("Excel 결과물 저장을 확인하지 못했습니다.")
                actual = target.Value
                if row_count == 1 and column_count == 1:
                    actual = ((actual,),)
                elif row_count == 1:
                    actual = (tuple(actual),)
                else:
                    actual = tuple(tuple(row) for row in actual)
                expected = tuple(tuple(row) for row in matrix)
                header_match = all(
                    _excel_cell_equal(actual[0][index], expected[0][index])
                    for index in range(column_count)
                )
                representative_match = (
                    _excel_cell_equal(actual[1][0], expected[1][0])
                    and _excel_cell_equal(actual[-1][-1], expected[-1][-1])
                )
                if (
                    len(actual) != row_count
                    or any(len(row) != column_count for row in actual)
                    or not header_match
                    or not representative_match
                ):
                    raise WorkflowError("Excel 표 읽기 검증에 실패했습니다.")
                return {
                    "path": str(destination),
                    "fingerprint": file_fingerprint(destination),
                    "verification": {
                        "exists": True,
                        "format": "xlsx",
                        "row_count": row_count - 1,
                        "column_count": column_count,
                        "header_match": True,
                        "representative_cells_match": True,
                    },
                }
            except Exception:
                if lease is not None:
                    lease.cleanup()
                    lease = None
                if created or destination.exists():
                    destination.unlink(missing_ok=True)
                raise
            finally:
                if lease is not None:
                    lease.cleanup()


class PdfOfficeWorkflow:
    """Prepare and atomically execute one allowlisted set of PDF artifacts."""

    def __init__(
        self,
        *,
        excel_writer=None,
        word_writer=None,
        hwp_writer=None,
        powerpoint_writer=None,
    ):
        self.excel_writer = excel_writer or PdfExcelTableWriter()
        self.word_writer = word_writer or WordReportWriter()
        self.hwp_writer = hwp_writer or HwpReportWriter()
        self.powerpoint_writer = powerpoint_writer or PowerPointSummaryWriter()

    def prepare(self, connection: PdfConnection, output_kinds) -> dict[str, Any]:
        if not isinstance(connection, PdfConnection):
            raise WorkflowError("PDF 결과물 준비에는 현재 연결 정보가 필요합니다.")
        outputs = tuple(str(item or "").strip().casefold() for item in output_kinds)
        recipe = pdf_workflow_step_recipe(outputs)
        source = Path(connection.file_path).resolve(strict=True)
        directory = source.parent
        stem = f"{_safe_stem(source.stem)}_PDF_분석"
        paths = {
            kind: str(_reserve_output(directory, stem, _OUTPUT_SUFFIXES[kind]))
            for kind in outputs
        }
        if "hwp" in outputs:
            self.hwp_writer.preflight()
        return {
            "output_kinds": list(outputs),
            "output_paths": paths,
            "output_names": [Path(paths[kind]).name for kind in outputs],
            "recipe": recipe,
        }

    @staticmethod
    def evidence(plan: Mapping[str, Any]) -> dict[str, Any]:
        outputs = tuple(plan.get("output_kinds") or ())
        return {
            "output_kinds": list(outputs),
            "output_names": [str(item) for item in plan.get("output_names") or ()],
            "recipe": validate_pdf_workflow_step_recipe(plan.get("recipe"), outputs),
        }

    @staticmethod
    def _validate_paths(connection: PdfConnection, plan: Mapping[str, Any]):
        outputs = tuple(plan.get("output_kinds") or ())
        validate_pdf_workflow_step_recipe(plan.get("recipe"), outputs)
        paths = dict(plan.get("output_paths") or {})
        if set(paths) != set(outputs):
            raise WorkflowError("PDF 결과물 계획의 파일 목록이 일치하지 않습니다.")
        source_parent = os.path.normcase(str(Path(connection.file_path).resolve().parent))
        normalized = {}
        for kind in outputs:
            path = Path(paths[kind]).resolve()
            if (
                kind not in _OUTPUT_SUFFIXES
                or path.suffix.casefold() != _OUTPUT_SUFFIXES[kind]
                or os.path.normcase(str(path.parent)) != source_parent
                or path.exists()
            ):
                raise WorkflowError("승인 후 PDF 결과물 경로가 바뀌었거나 이미 존재합니다.")
            normalized[kind] = path
        return outputs, normalized

    @staticmethod
    def _product(connection, answer: str, citation_pages) -> WorkProductData:
        citation = citation_label(citation_pages)
        return WorkProductData(
            title=f"{Path(connection.document_name).stem} PDF 분석",
            metrics=[],
            tables=[],
            charts=[],
            insights=[str(answer), citation],
            source_files=[connection.file_path],
        )

    @staticmethod
    def _verify_reopened(kind: str, path: Path, citation: str) -> dict[str, Any]:
        if kind == "word":
            from docx import Document

            reopened = Document(str(path))
            text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
            if citation not in text:
                raise WorkflowError("Word 결과물 재열기 인용 검증에 실패했습니다.")
            return {"reopened": True, "citation_present": True}
        if kind == "powerpoint":
            from pptx import Presentation

            reopened = Presentation(str(path))
            text = "\n".join(
                str(shape.text)
                for slide in reopened.slides
                for shape in slide.shapes
                if hasattr(shape, "text")
            )
            if citation not in text:
                raise WorkflowError("PowerPoint 결과물 재열기 인용 검증에 실패했습니다.")
            return {
                "reopened": True,
                "citation_present": True,
                "slide_count": len(reopened.slides),
            }
        return {"reopened": False, "citation_present": True}

    def execute(
        self,
        connection: PdfConnection,
        plan: Mapping[str, Any],
        *,
        answer: str = "",
        citation_pages=(),
        table: PdfTableCandidate | None = None,
    ) -> dict[str, Any]:
        outputs, paths = self._validate_paths(connection, plan)
        citation = citation_label(citation_pages) if citation_pages else ""
        product = None
        if any(kind in outputs for kind in ("word", "hwp", "powerpoint")):
            if not str(answer).strip() or not citation_pages:
                raise WorkflowError("보고서 생성에는 페이지 인용이 있는 PDF 분석 결과가 필요합니다.")
            product = self._product(connection, answer, citation_pages)
        results = {}
        created: list[Path] = []
        try:
            for kind in outputs:
                path = paths[kind]
                if kind == "excel":
                    result = self.excel_writer.run(path, table)
                else:
                    context = {
                        "output_path": str(path),
                        "preferences": {"_source_kind": "pdf"},
                        "report_format": "hwp" if kind == "hwp" else "word",
                        "slide_count": 5,
                    }
                    if kind == "word":
                        result = self.word_writer.run(context, product)
                    elif kind == "hwp":
                        result = self.hwp_writer.run(context, product)
                    else:
                        result = self.powerpoint_writer.run(context, product)
                    result = dict(result)
                    verification = dict(result.get("verification") or {})
                    if kind == "hwp":
                        if not (
                            verification.get("reopened") is True
                            and verification.get("citation_present") is True
                        ):
                            raise WorkflowError(
                                "한글 PDF 보고서 재열기 검증이 확인되지 않았습니다."
                            )
                    else:
                        verification.update(
                            self._verify_reopened(kind, path, citation)
                        )
                    result["verification"] = verification
                created.append(path)
                results[kind] = result
            return {
                "outputs": results,
                "output_names": [paths[kind].name for kind in outputs],
                "recipe": validate_pdf_workflow_step_recipe(
                    plan.get("recipe"), outputs
                ),
            }
        except Exception:
            for path in created:
                path.unlink(missing_ok=True)
            raise


__all__ = ["PdfExcelTableWriter", "PdfOfficeWorkflow"]
