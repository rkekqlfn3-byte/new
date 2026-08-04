import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation

from engine.pdf import PdfIntakeManager, PdfOfficeWorkflow, find_pdf_tables
from engine.pdf.reader import read_pdf_document
from tests.fixtures.pdf_factory import create_owned_pdf_fixtures


def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class FakeArtifactWriter:
    def __init__(self, kind, *, fail=False):
        self.kind = kind
        self.fail = fail
        self.calls = []
        self.preflight_calls = 0

    def preflight(self):
        self.preflight_calls += 1
        return {"status": "ready"}

    def run(self, context, product):
        self.calls.append((context, product))
        if self.fail:
            raise RuntimeError("owned writer failure")
        path = Path(context["output_path"])
        citation = next(
            item for item in product.insights if str(item).startswith("근거 페이지:")
        )
        if self.kind == "word":
            document = Document()
            document.add_paragraph(product.title)
            document.add_paragraph(citation)
            document.save(path)
        elif self.kind == "powerpoint":
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = product.title
            slide.placeholders[1].text = citation
            presentation.save(path)
        else:
            path.write_text(f"{product.title}\n{citation}", encoding="utf-8")
        return {
            "path": str(path),
            "fingerprint": {"sha256": fingerprint(path)},
            "verification": {
                "exists": True,
                "content_readback": True,
                "reopened": self.kind == "hwp",
                "citation_present": True,
            },
        }


class FakeExcelWriter:
    def __init__(self):
        self.calls = []

    def run(self, output_path, table):
        self.calls.append((Path(output_path), table))
        path = Path(output_path)
        path.write_text("owned xlsx placeholder", encoding="utf-8")
        return {
            "path": str(path),
            "fingerprint": {"sha256": fingerprint(path)},
            "verification": {
                "exists": True,
                "row_count": len(table.rows),
                "column_count": len(table.headers),
                "header_match": True,
                "representative_cells_match": True,
            },
        }


class PdfOfficeWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="jarvis-pdf-office-")
        self.addCleanup(self.temporary.cleanup)
        self.fixtures = create_owned_pdf_fixtures(self.temporary.name)
        self.manager = PdfIntakeManager()
        self.manager.connect_file(str(self.fixtures.text))
        self.connection = self.manager.current()

    def test_report_artifacts_are_new_reopened_and_page_grounded(self):
        word = FakeArtifactWriter("word")
        hwp = FakeArtifactWriter("hwp")
        powerpoint = FakeArtifactWriter("powerpoint")
        workflow = PdfOfficeWorkflow(
            word_writer=word,
            hwp_writer=hwp,
            powerpoint_writer=powerpoint,
        )
        source_before = fingerprint(self.fixtures.text)

        plan = workflow.prepare(
            self.connection,
            ("word", "hwp", "powerpoint"),
        )
        evidence = repr(workflow.evidence(plan))
        result = workflow.execute(
            self.connection,
            plan,
            answer="Owned grounded answer [p.1]",
            citation_pages=(1,),
        )

        self.assertEqual(source_before, fingerprint(self.fixtures.text))
        self.assertEqual(1, hwp.preflight_calls)
        self.assertNotIn(self.temporary.name, evidence)
        self.assertEqual(3, len(result["outputs"]))
        self.assertTrue(result["outputs"]["word"]["verification"]["reopened"])
        self.assertTrue(
            result["outputs"]["powerpoint"]["verification"]["citation_present"]
        )
        self.assertTrue(all(Path(item["path"]).is_file() for item in result["outputs"].values()))

    def test_excel_table_output_uses_verified_candidate_without_ai(self):
        self.manager.connect_file(str(self.fixtures.table))
        connection = self.manager.current()
        table = find_pdf_tables(read_pdf_document(str(self.fixtures.table)))[0]
        excel = FakeExcelWriter()
        workflow = PdfOfficeWorkflow(excel_writer=excel)

        plan = workflow.prepare(connection, ("excel",))
        result = workflow.execute(
            connection,
            plan,
            table=table,
            citation_pages=(table.page_number,),
        )

        verification = result["outputs"]["excel"]["verification"]
        self.assertTrue(verification["header_match"])
        self.assertTrue(verification["representative_cells_match"])
        self.assertEqual(1, len(excel.calls))

    def test_later_writer_failure_removes_earlier_artifacts(self):
        word = FakeArtifactWriter("word")
        powerpoint = FakeArtifactWriter("powerpoint", fail=True)
        workflow = PdfOfficeWorkflow(
            word_writer=word,
            powerpoint_writer=powerpoint,
        )
        plan = workflow.prepare(self.connection, ("word", "powerpoint"))

        with self.assertRaisesRegex(RuntimeError, "owned writer failure"):
            workflow.execute(
                self.connection,
                plan,
                answer="Owned answer [p.1]",
                citation_pages=(1,),
            )

        self.assertTrue(all(not Path(path).exists() for path in plan["output_paths"].values()))


if __name__ == "__main__":
    unittest.main()
